#!/usr/bin/env python3
"""
Anchored Index Builder - Builds anchored inverted indexes from Excel files using OpenAI analysis.
Includes cell references (A1, D3) with values for precise data location.
Classes: OpenAIClient, ExcelProcessor, ChunkAnalyzer, SheetProcessor, AnchoredIndexBuilder
Functions: process_sheet (with overlap), build_index, query_index, save_index
"""

import openpyxl
import json
import os
from typing import Dict, List, Any, Union
from datetime import datetime
from collections import defaultdict
import re
import requests
import time
from dotenv import load_dotenv

load_dotenv()

EXCEL_FILE_PATH = os.getenv("EXCEL_FILE_PATH", "test.xlsx")
CHUNK_SIZE = int(os.getenv("CHUNK_SIZE", "50"))
CHUNK_OVERLAP_PERCENT = float(os.getenv("CHUNK_OVERLAP_PERCENT", "5.0"))
CHUNK_OVERLAP_SIZE = max(1, int(CHUNK_SIZE * CHUNK_OVERLAP_PERCENT / 100))


class OpenAIClient:
    def __init__(self, api_key: str = None, model: str = "gpt-4o-mini"):
        self.api_key = api_key or os.getenv("OPENAI_API_KEY")
        self.model = model
        if not self.api_key:
            raise ValueError("OpenAI API key not provided")
    
    def call_api(self, prompt: str, max_retries: int = 3) -> str:
        for attempt in range(max_retries):
            try:
                headers = {
                    "Authorization": f"Bearer {self.api_key}",
                    "Content-Type": "application/json"
                }
                
                data = {
                    "model": self.model,
                    "messages": [{"role": "user", "content": prompt}],
                    "max_tokens": 500,
                    "temperature": 0.0
                }
                
                response = requests.post(
                    "https://api.openai.com/v1/chat/completions",
                    headers=headers,
                    json=data,
                    timeout=30
                )
                
                if response.status_code == 200:
                    return response.json()["choices"][0]["message"]["content"].strip()
                else:
                    raise Exception(f"API call failed: {response.status_code}")
                    
            except Exception as e:
                if attempt == max_retries - 1:
                    raise e
                time.sleep(2 ** attempt)
        
        raise Exception("Max retries exceeded")


class ExcelProcessor:
    def __init__(self, file_path: str = None):
        self.file_path = file_path or EXCEL_FILE_PATH
        self.formula_workbook = None
    
    def _get_formula_workbook(self):
        if self.formula_workbook is None:
            self.formula_workbook = openpyxl.load_workbook(self.file_path, data_only=False)
        return self.formula_workbook
    
    def extract_headers(self, sheet) -> List[str]:
        headers = []
        for col in range(1, sheet.max_column + 1):
            cell_value = sheet.cell(row=1, column=col).value
            if cell_value is not None:
                headers.append(str(cell_value).strip())
            else:
                headers.append(f"Column_{col}")
        return headers
    
    def extract_chunk_data(self, sheet, start_row: int, end_row: int, max_col: int, headers: List[str]) -> tuple:
        chunk_data = []
        formulas_data = []
        
        # Get the formula workbook for formula extraction
        formula_wb = self._get_formula_workbook()
        formula_sheet = formula_wb[sheet.title] if sheet.title in formula_wb.sheetnames else formula_wb.active
        
        for row_idx in range(start_row, end_row + 1):
            row_data = []
            row_formulas = []
            
            for col_idx in range(1, max_col + 1):
                # Use original sheet for cell values
                cell = sheet.cell(row=row_idx, column=col_idx)
                cell_value = cell.value
                cell_coordinate = cell.coordinate
                
                # Use formula sheet for formula detection
                formula_cell = formula_sheet.cell(row=row_idx, column=col_idx)
                
                # Always check for formulas first
                if formula_cell.data_type == 'f':
                    formula_ref = f"{formula_cell.coordinate}"
                    row_formulas.append(f"{formula_ref}: {formula_cell.value}")
                
                # Then check for cell values (formulas might have None values in data_only=True mode)
                if cell_value is not None:
                    if isinstance(cell_value, str):
                        cell_str = cell_value.strip()
                    else:
                        cell_str = str(cell_value)
                    
                    if cell_str:
                        # Include cell reference with value: "A1: value"
                        cell_with_ref = f"{cell_coordinate}: {cell_str}"
                        row_data.append(cell_with_ref)
                elif formula_cell.data_type == 'f':
                    # For formula cells with no value, include the formula as the value
                    cell_with_ref = f"{cell_coordinate}: {formula_cell.value}"
                    row_data.append(cell_with_ref)
            
            if row_data:
                row_string = " | ".join(row_data)
                chunk_data.append(row_string)
            
            if row_formulas:
                formulas_data.append(f"Row{row_idx}: {'; '.join(row_formulas)}")
        
        return chunk_data, formulas_data


class ChunkAnalyzer:
    def __init__(self, openai_client: OpenAIClient):
        self.openai_client = openai_client
    
    def analyze_chunk(self, chunk_data: List[str], chunk_range: str) -> Dict[str, str]:
        if not chunk_data:
            return {"summary": "Empty chunk", "context": "No data available"}
        
        data_text = "\n".join(chunk_data)
        
        prompt = f"""
Analyze this Excel data chunk (rows {chunk_range}) and provide:

1. SUMMARY: A concise 4-6 sentence summary of what this data represents
2. CONTEXT: Key details, patterns, or important information in this data. Be specific and detailed. not more that 5 lines.

Data:
{data_text}

Respond in this exact format:
SUMMARY: [your summary here]
CONTEXT: [your context analysis here]
"""
        
        try:
            response = self.openai_client.call_api(prompt)
            
            summary_match = re.search(r'SUMMARY:\s*(.+)', response, re.IGNORECASE | re.DOTALL)
            context_match = re.search(r'CONTEXT:\s*(.+)', response, re.IGNORECASE | re.DOTALL)
            
            summary = summary_match.group(1).strip() if summary_match else "Data analysis unavailable"
            context = context_match.group(1).strip() if context_match else "Context analysis unavailable"
            
            return {"summary": summary, "context": context}
            
        except Exception as e:
            print(f"OpenAI analysis failed: {e}")
            return {
                "summary": f"Data chunk from {chunk_range}",
                "context": f"Raw data: {len(chunk_data)} rows"
            }


class SheetProcessor:
    def __init__(self, excel_processor: ExcelProcessor, chunk_analyzer: ChunkAnalyzer):
        self.excel_processor = excel_processor
        self.chunk_analyzer = chunk_analyzer
    
    def process_sheet(self, sheet, chunk_size: int = None) -> Dict:
        if chunk_size is None:
            chunk_size = CHUNK_SIZE
            
        max_row = sheet.max_row
        max_col = sheet.max_column
        
        headers = self.excel_processor.extract_headers(sheet)
        
        index = {
            "sheet_name": sheet.title,
            "num_rows": max_row,
            "num_cols": max_col,
            "headers": headers,
            "chunk_size": chunk_size,
            "chunk_overlap_size": CHUNK_OVERLAP_SIZE,
            "anchors": {},
            "metadata": {
                "total_chunks": 0,
                "processing_time": 0,
                "openai_calls": 0
            }
        }
        
        start_time = time.time()
        chunk_num = 0
        start_row = 1
        max_chunks = max_row + 10  # Safety limit to prevent infinite loops
        
        while start_row <= max_row and chunk_num < max_chunks:
            end_row = min(start_row + chunk_size - 1, max_row)
            chunk_range = f"A{start_row}:A{end_row}"
            
            chunk_data, formulas_data = self.excel_processor.extract_chunk_data(
                sheet, start_row, end_row, max_col, headers
            )
            
            try:
                print(f"Analyzing chunk {chunk_num + 1}: {chunk_range}")
                analysis = self.chunk_analyzer.analyze_chunk(chunk_data, chunk_range)
                index["metadata"]["openai_calls"] += 1
                
                anchor_key = f"chunk_{chunk_num}"
                index["anchors"][anchor_key] = {
                    "range": chunk_range,
                    "start_row": start_row,
                    "end_row": end_row,
                    "summary": analysis["summary"],
                    "context": analysis["context"],
                    "table": chunk_data,
                    "formulas": formulas_data,
                    "chunk_number": chunk_num
                }
                
                chunk_num += 1
                
                # Calculate next start row with overlap
                next_start_row = end_row - CHUNK_OVERLAP_SIZE + 1
                
                # If next start row would be beyond max_row or same as current start_row, we're done
                if next_start_row > max_row or next_start_row <= start_row:
                    break
                
                start_row = next_start_row
                
            except Exception as e:
                print(f"Failed to analyze chunk {chunk_range}: {e}")
                raise e
        
        index["metadata"]["total_chunks"] = chunk_num
        index["metadata"]["processing_time"] = time.time() - start_time
        
        return index


class AnchoredIndexBuilder:
    def __init__(self, openai_api_key: str = None, openai_model: str = "gpt-4o-mini", file_path: str = None):
        self.file_path = file_path or EXCEL_FILE_PATH
        self.openai_client = OpenAIClient(api_key=openai_api_key, model=openai_model)
        self.excel_processor = ExcelProcessor(self.file_path)
        self.chunk_analyzer = ChunkAnalyzer(self.openai_client)
        self.sheet_processor = SheetProcessor(self.excel_processor, self.chunk_analyzer)
    
    def build_index(self, file_path: str = None, sheet_name: Union[int, str, None] = None, chunk_size: int = None) -> Dict:
        if file_path is None:
            file_path = EXCEL_FILE_PATH
        if chunk_size is None:
            chunk_size = CHUNK_SIZE
            
        wb = openpyxl.load_workbook(file_path, data_only=True)
        
        if sheet_name is None:
            sheets_to_process = wb.worksheets
        elif isinstance(sheet_name, int):
            sheets_to_process = [wb.worksheets[sheet_name]]
        elif isinstance(sheet_name, str):
            sheets_to_process = [wb[sheet_name]]
        else:
            sheets_to_process = [wb.active]
        
        all_sheets_index = {"sheets": {}}
        
        for sheet in sheets_to_process:
            print(f"Processing sheet: {sheet.title}")
            sheet_index = self.sheet_processor.process_sheet(sheet, chunk_size)
            all_sheets_index["sheets"][sheet.title] = sheet_index
        
        return all_sheets_index
    
    def query_index(self, index: Dict, query_terms: List[str]) -> Dict:
        results = {}
        
        for sheet_name, sheet_data in index["sheets"].items():
            sheet_results = {
                "sheet_info": {
                    "name": sheet_name,
                    "total_chunks": sheet_data["metadata"]["total_chunks"],
                    "headers": sheet_data["headers"]
                },
                "matching_chunks": []
            }
            
            for chunk_id, chunk_data in sheet_data["anchors"].items():
                chunk_text = f"{chunk_data['summary']} {chunk_data['context']} {' '.join(chunk_data['table'])}"
                
                for term in query_terms:
                    if term.lower() in chunk_text.lower():
                        sheet_results["matching_chunks"].append({
                            "chunk_id": chunk_id,
                            "range": chunk_data["range"],
                            "summary": chunk_data["summary"],
                            "context": chunk_data["context"],
                            "table": chunk_data["table"][:3],
                            "match_term": term
                        })
                        break
            
            if sheet_results["matching_chunks"]:
                results[sheet_name] = sheet_results
        
        return results
    
    def save_index(self, index: Dict, file_path: str = None):
        if file_path is None:
            file_path = os.getenv("ANCHORED_INDEX_OUTPUT_PATH", "anchored_index_output.json")
        with open(file_path, "w") as f:
            json.dump(index, f, indent=2, default=str)
        print(f"Index saved to '{file_path}'")


def main():
    try:
        print("Building anchored inverted index...")
        print(f"Excel file: {EXCEL_FILE_PATH}")
        print(f"Chunk size: {CHUNK_SIZE} rows")
        print(f"Overlap: {CHUNK_OVERLAP_SIZE} rows ({CHUNK_OVERLAP_PERCENT}%)")
        print("This may take a few minutes...")
        
        builder = AnchoredIndexBuilder()
        index = builder.build_index()
        
        print(f"Anchored index built successfully!")
        print(f"Found {len(index['sheets'])} sheet(s)")
        
        for sheet_name, sheet_data in index["sheets"].items():
            print(f"Sheet '{sheet_name}': {sheet_data['num_rows']} rows, {sheet_data['num_cols']} cols")
            print(f"  Chunks: {sheet_data['metadata']['total_chunks']}")
            print(f"  Processing time: {sheet_data['metadata']['processing_time']:.2f}s")
        
        builder.save_index(index)
        
        print("\n" + "="*50)
        
        # Automatically upload to Pinecone after building the index
        print("\n🚀 Auto-uploading to Pinecone...")
        try:
            from pinecone_uploader_new import PineconeUploader, load_anchored_index
            
            # Load the just-created anchored index
            index_data = load_anchored_index()
            
            # Upload to Pinecone
            uploader = PineconeUploader()
            uploader.upload_data(index_data, embedding_method="openai")
            
            print("✅ Automatic Pinecone upload completed!")
            
        except Exception as upload_error:
            print(f"⚠️ Auto-upload to Pinecone failed: {upload_error}")
            print("You can manually run: python3 pinecone_uploader_new.py")
       
        
    except Exception as e:
        print(f"Error: {e}")
        import traceback
        traceback.print_exc()


if __name__ == "__main__":
    main()
